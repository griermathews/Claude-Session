from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Inches(0.85)
section.bottom_margin = Inches(0.85)
section.left_margin   = Inches(1.0)
section.right_margin  = Inches(1.0)

# ── Colour palette ────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1A, 0x52, 0x76)   # headings
RED    = RGBColor(0xC0, 0x39, 0x2B)   # emphasis / risk
GREY   = RGBColor(0x55, 0x55, 0x55)   # body
BLACK  = RGBColor(0x00, 0x00, 0x00)

def set_color(run, rgb):
    run.font.color.rgb = rgb

def heading(text, size=13, color=NAVY, bold=True, space_before=14, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_color(run, color)
    return p

def body(text, size=10.5, color=GREY, space_before=2, space_after=2,
         italic=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    run.bold   = bold
    set_color(run, color)
    return p

def bullet(text, size=10.5, color=GREY, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.size = Pt(size)
        set_color(rb, BLACK)
        rt = p.add_run(text)
        rt.font.size = Pt(size)
        set_color(rt, color)
    else:
        r = p.add_run(text)
        r.font.size = Pt(size)
        set_color(r, color)
    return p

def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.append(pBdr)

# ═══════════════════════════════════════════════════════════════════════════════
# HEADER BLOCK
# ═══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(2)
r = p.add_run("MERIDIAN TECHNOLOGIES")
r.bold = True; r.font.size = Pt(9)
set_color(r, GREY)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after  = Pt(2)
r2 = p2.add_run("Board of Directors — Annual Strategic Review")
r2.font.size = Pt(9)
set_color(r2, GREY)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_before = Pt(0)
p3.paragraph_format.space_after  = Pt(6)
r3 = p3.add_run("CEO Opening Remarks  |  Catherine Park")
r3.font.size = Pt(9)
set_color(r3, GREY)

divider()

# ── Headline ──────────────────────────────────────────────────────────────────
p_hl = doc.add_paragraph()
p_hl.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_hl.paragraph_format.space_before = Pt(10)
p_hl.paragraph_format.space_after  = Pt(10)
r_hl = p_hl.add_run(
    "Meridian is profitable, decelerating, and at an inflection point:\n"
    "the decisions we make in 2026 will determine whether we are a managed-decline\n"
    "business or the defining agentic-work platform in regulated enterprise."
)
r_hl.bold = True
r_hl.font.size = Pt(12)
set_color(r_hl, NAVY)

divider()

# ═══════════════════════════════════════════════════════════════════════════════
# CHART
# ═══════════════════════════════════════════════════════════════════════════════
heading("The Numbers in One Picture", size=11, space_before=10, space_after=4)

body(
    "Revenue growth has decelerated every year since 2022 — 28% → 19% → 16% → 11%. "
    "Operating and free-cash-flow margins have expanded every year. "
    "This is what a business looks like when cost discipline outpaces growth investment. "
    "Both trends are intentional. Neither is sustainable indefinitely.",
    size=10, space_before=0, space_after=6
)

doc.add_picture("key_trend_chart.png", width=Inches(5.8))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(3)
cap.paragraph_format.space_after  = Pt(2)
rc = cap.add_run("Figure 1 — Revenue growth (bars, left axis) vs. operating and FCF margins (lines, right axis), 2023–2025.")
rc.italic = True; rc.font.size = Pt(8.5)
set_color(rc, GREY)

divider()

# ═══════════════════════════════════════════════════════════════════════════════
# THREE ISSUES
# ═══════════════════════════════════════════════════════════════════════════════
heading("Three Issues for the Board", size=13, space_before=12)

# ── Issue 1 ───────────────────────────────────────────────────────────────────
heading("1.  Growth deceleration is structural — and the 2026 plan does not fully reverse it.",
        size=11, color=RED, space_before=10, space_after=3)

body(
    "We will guide 2026 at 10–14% growth. The plan we approved in 2022 assumed $560M "
    "revenue by 2026; we will deliver $440–455M — a $100M+ miss. This is not a "
    "communication problem. It is a structural problem.",
    size=10, space_before=0, space_after=4
)

bullet("Revenue growth: 28% (2022) → 19% (2023) → 16% (2024) → 11% (2025). "
       "No quarter has reversed the trend.",
       bold_prefix="Trend: ")
bullet("Magic number fell from 1.20 (Q1 2024) to 0.92 (Q4 2025). "
       "CAC payback stretched from 18 to 22 months. "
       "We are spending more per dollar of new ARR, not less.",
       bold_prefix="Efficiency: ")
bullet("SMB, once 38% of ARR, is now 13% and declining 23% YoY. "
       "The prior plan called this a controlled wind-down. It was a hemorrhage.",
       bold_prefix="Mix: ")

body(
    "The board needs to understand: reacceleration requires a second leg beyond "
    "enterprise growth. My stress-test shows Copilot at per-seat pricing cannot "
    "close the gap alone. Consumption pricing — currently a Q3 2026 roadmap item "
    "— is the single most leveraged variable in our 2027 growth rate.",
    size=10, italic=True, space_before=5, space_after=2
)

# ── Issue 2 ───────────────────────────────────────────────────────────────────
heading("2.  The AI Copilot bet is real — and fragile.",
        size=11, color=RED, space_before=10, space_after=3)

body(
    "Copilot reached GA in September 2025, hit 710 paying seats by year-end "
    "against an internal target of 600, and achieved a 44% enterprise renewal "
    "attach rate. The Helio acquisition has already compressed our agent "
    "architecture roadmap by an estimated 15 months. These are genuine wins. "
    "Three risks sit directly underneath them.",
    size=10, space_before=0, space_after=4
)

bullet("Asana shipped its full agent suite in November. Monday now has larger ARR "
       "than Meridian. ClearAI Work raised $120M at $1B+ valuation. "
       "Atlassian's agentic Jira will appear in our enterprise deals directly. "
       "The window to establish Copilot as the enterprise-grade choice is 2026.",
       bold_prefix="Competitive: ")
bullet("The 2026 roadmap requires 80 net new engineers in H1. "
       "At 15% annual attrition, actual net adds are closer to 25. "
       "Six of twelve 2025 roadmap items slipped at least one quarter.",
       bold_prefix="Capacity: ")
bullet("The Helio retention equity vests over four years with a cash cliff in 2026. "
       "26 of 28 engineers are still here. The agent-builder roadmap (Q2 2026) "
       "depends on keeping them.",
       bold_prefix="Retention: ")

body(
    "The board should ask: what is the specific plan to close the engineering "
    "capacity gap, and what is the contingency if Helio retention slips post-cliff?",
    size=10, italic=True, space_before=5, space_after=2
)

# ── Issue 3 ───────────────────────────────────────────────────────────────────
heading("3.  Mid-market is the quiet risk no one is talking about.",
        size=11, color=RED, space_before=10, space_after=3)

body(
    "Enterprise is healthy. SMB decline is visible and being managed. "
    "Mid-market — 47% of our ARR — is quietly eroding and has not found "
    "a floor.",
    size=10, space_before=0, space_after=4
)

bullet("Mid-market NRR has declined every single quarter in our dataset: "
       "108% (Q1 2024) → 102% (Q4 2025). Gross logo churn rose from 9.1% to 10.2%.",
       bold_prefix="NRR compression: ")
bullet("Asana is now bundling AI into its standard tier. "
       "Mid-market renewals are pushing for price holds, expanded user counts at flat "
       "price, or Copilot bundled at no cost. We are holding the line — for now.",
       bold_prefix="Pricing pressure: ")
bullet("The resource management module — the top-3 ask from our customer advisory board "
       "and the primary mid-market expansion lever — was deferred in 2025 and is "
       "not scheduled until Q2 2026.",
       bold_prefix="Product gap: ")
bullet("28% of sales respondents in the 2025 employee survey said they "
       "cannot crisply differentiate Meridian from Asana with AI.",
       bold_prefix="Sales confidence: ")

body(
    "If mid-market NRR continues compressing into the high 90s, this segment becomes "
    "a retention story. The math on 15%+ company growth becomes very hard even with "
    "strong enterprise execution.",
    size=10, italic=True, space_before=5, space_after=2
)

divider()

# ═══════════════════════════════════════════════════════════════════════════════
# ONE ASK
# ═══════════════════════════════════════════════════════════════════════════════
heading("My One Ask of the Board", size=13, space_before=12, space_after=4)

body(
    "I am not asking for approval of a full three-year plan today. "
    "I will bring that to the board no later than Q1 2026, consistent with the "
    "authorization given at the time of my appointment. "
    "What I am asking for today is one decision:",
    size=10, space_before=0, space_after=6
)

p_ask = doc.add_paragraph()
p_ask.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ask.paragraph_format.space_before = Pt(6)
p_ask.paragraph_format.space_after  = Pt(6)
p_ask.paragraph_format.left_indent  = Inches(0.5)
p_ask.paragraph_format.right_indent = Inches(0.5)
r_ask = p_ask.add_run(
    "Authorize management to move the Copilot consumption-pricing model "
    "from a Q3 2026 roadmap item to a Q1 2026 strategic priority, "
    "with CFO, CRO, and CPO jointly accountable for a board-ready proposal "
    "by the March meeting."
)
r_ask.bold = True
r_ask.font.size = Pt(11)
set_color(r_ask, NAVY)

body(
    "This is the single decision with the highest leverage on our 2027 growth rate. "
    "It does not require additional capital. It requires alignment at the executive "
    "level and a clear signal from this board that growth reacceleration — not margin "
    "expansion — is the primary objective for 2026.",
    size=10, space_before=6, space_after=4
)

divider()

# ── Footer note ───────────────────────────────────────────────────────────────
p_foot = doc.add_paragraph()
p_foot.paragraph_format.space_before = Pt(8)
r_foot = p_foot.add_run(
    "Supporting materials: meridian_financials_2022_2025.csv  |  meridian_kpis_2024.csv  |  "
    "meridian_segments_overview.md  |  meridian_product_roadmap_2025.md  |  "
    "meridian_earnings_call_q1–q4_2025.txt  |  copilot_stress_test.png"
)
r_foot.font.size = Pt(7.5)
r_foot.italic = True
set_color(r_foot, GREY)

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save("board_opening_remarks.docx")
print("Saved board_opening_remarks.docx")
